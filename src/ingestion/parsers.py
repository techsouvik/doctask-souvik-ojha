"""Document parsers for DOCX, PDF, and TXT files."""

import os
from typing import Tuple, List, Dict, Any
import fitz  # PyMuPDF
from docx import Document as DocxDocument


def parse_docx(filepath: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse a .docx file and return full text and structural blocks."""
    doc = DocxDocument(filepath)
    full_text_parts = []
    blocks = []
    
    # Process paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        full_text_parts.append(text)
        style_name = p.style.name if (p.style and p.style.name) else "Normal"
        is_heading = style_name.startswith("Heading") or style_name.startswith("Title")
        
        blocks.append({
            "type": "heading" if is_heading else "paragraph",
            "title": text if is_heading else None,
            "text": text,
            "location": f"Paragraph {p_idx + 1}" + (f" ({style_name})" if is_heading else "")
        })

    # Process tables
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        
        if table_rows:
            table_text = "\n".join(table_rows)
            full_text_parts.append(f"\n[TABLE {t_idx + 1}]\n" + table_text)
            blocks.append({
                "type": "table",
                "title": f"Table {t_idx + 1}",
                "text": table_text,
                "location": f"Table {t_idx + 1}"
            })

    return "\n\n".join(full_text_parts), blocks


def parse_pdf(filepath: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse a .pdf file and return full text and structural blocks with page numbers."""
    doc = fitz.open(filepath)
    full_text_parts = []
    blocks = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        raw_text = page.get_text()
        text = str(raw_text).strip() if raw_text else ""
        if not text:
            continue
        
        full_text_parts.append(f"--- Page {page_num + 1} ---\n{text}")
        blocks.append({
            "type": "page_block",
            "title": f"Page {page_num + 1}",
            "text": text,
            "page_num": page_num + 1,
            "location": f"Page {page_num + 1}"
        })

    doc.close()
    return "\n\n".join(full_text_parts), blocks


def parse_txt(filepath: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse a .txt file and return full text and structural blocks."""
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        text = f.read().strip()

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    blocks = []
    for idx, p in enumerate(paragraphs):
        blocks.append({
            "type": "paragraph",
            "title": None,
            "text": p,
            "location": f"Section/Paragraph {idx + 1}"
        })

    return text, blocks


def parse_document(filepath: str) -> Tuple[str, List[Dict[str, Any]], str]:
    """Auto-detect format and parse file."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".docx":
        text, blocks = parse_docx(filepath)
        fmt = "docx"
    elif ext == ".pdf":
        text, blocks = parse_pdf(filepath)
        fmt = "pdf"
    elif ext in [".txt", ".md"]:
        text, blocks = parse_txt(filepath)
        fmt = "txt"
    else:
        text, blocks = parse_txt(filepath)
        fmt = "raw"
        
    return text, blocks, fmt
